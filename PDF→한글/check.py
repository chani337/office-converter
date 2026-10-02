"""설치된 한글에서 실제 DOCX→HWP 저장/다시 열기 검사."""
from pathlib import Path
import json
from docx import Document
from docx.oxml.ns import qn
from converter import Converter

root=Path(__file__).resolve().parent/'호환성 점검'
root.mkdir(exist_ok=True)
doc=Document()
doc.styles['Normal'].font.name='맑은 고딕'
p=doc.add_paragraph('한글 2018 / 2022 호환성 점검입니다. 금액 123,456원.')
p.runs[0]._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'맑은 고딕')
table=doc.add_table(rows=2,cols=2)
table.style='Table Grid'
for cell,text in zip([c for row in table.rows for c in row.cells],['항목','금액','교재','10000']):cell.text=text
docx=root/'호환성확인.docx';hwp=root/'호환성확인.hwp'
doc.save(docx)
try:
    with Converter({'convert_timeout_sec':180}) as engine:
        report=engine.convert(docx,hwp)
    (root/'결과.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    if report['hwp_text_retention']<0.99:
        print('글자 보존 검사에서 차이가 있습니다. 호환성확인.hwp를 직접 열어 확인하세요.')
        raise SystemExit(2)
    print('저장/다시 열기/글자 보존 검사 통과. 실제 설치된 한글 버전:',report['hancom_version'])
    print('표와 배치는 호환성확인.hwp를 열어서 확인하세요.')
except Exception as exc:
    print('호환성 점검 실패:',exc)
    raise SystemExit(1)
