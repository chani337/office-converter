import sys
import tempfile
import unittest
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx_compat import normalize_docx, repair_grid

class LayoutTests(unittest.TestCase):
    def test_merged_cells_determine_grid_and_continuation_width(self):
        d=Document();t=d.add_table(rows=3, cols=3)
        for row in t.rows:
            for c,w in zip(row.cells,(1000,2000,3000)):
                c._tc.tcPr.tcW.w=w
        t.cell(0,0).merge(t.cell(1,0))
        t.cell(2,1).merge(t.cell(2,2))
        t._tbl.tr_lst[1].tc_lst[0].tcPr.tcW.w=2400
        self.assertTrue(repair_grid(t._tbl))
        self.assertEqual([int(c.get(qn('w:w'))) for c in t._tbl.tblGrid], [1000,2000,3000])
        self.assertEqual(t._tbl.tr_lst[1].tc_lst[0].tcPr.tcW.w,1000)
        self.assertEqual(t.cell(2,1)._tc.tcPr.tcW.w,5000)

    def test_conflicting_widths_not_guessed(self):
        d=Document();t=d.add_table(rows=2, cols=2)
        t.cell(1,0)._tc.tcPr.tcW.w=500
        before=t._tbl.xml
        self.assertFalse(repair_grid(t._tbl))
        self.assertEqual(before,t._tbl.xml)

    def test_fonts_borders_and_valid_numbers(self):
        with tempfile.TemporaryDirectory() as tmp:
            p=Path(tmp)/'test.docx';d=Document()
            for name in ('Batang','Arial','ABCDEF+MalgunGothic'):
                d.add_paragraph().add_run('한글 text').font.name=name
            t=d.add_table(rows=1,cols=1)
            borders=OxmlElement('w:tcBorders')
            for name in ('start','end','top','bottom'):
                b=OxmlElement('w:'+name)
                for k,v in {'sz':'2.8800001','color':'#000000','val':'single'}.items():b.set(qn('w:'+k),v)
                borders.append(b)
            t.cell(0,0)._tc.tcPr.append(borders);d.save(p)
            normalize_docx(p);result=Document(p)
            self.assertEqual([r.font.name for par in result.paragraphs for r in par.runs],['바탕','Arial','맑은 고딕'])
            self.assertFalse(result.element.xpath('.//w:tcBorders/w:start | .//w:tcBorders/w:end'))
            self.assertTrue(result.element.xpath('.//w:tcBorders/w:left | .//w:tcBorders/w:right'))
            for b in result.element.xpath('.//w:tcBorders/*'):
                self.assertEqual(b.get(qn('w:sz')),'3');self.assertEqual(b.get(qn('w:color')),'000000')
            before=result.element.xml;normalize_docx(p)
            self.assertEqual(Document(p).element.xml,before)

if __name__=='__main__':unittest.main()
