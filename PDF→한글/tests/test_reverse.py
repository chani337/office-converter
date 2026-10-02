import json
from pathlib import Path
import shutil
import sys
import tempfile
import types
import unittest
from unittest.mock import MagicMock,patch
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
import pymupdf
from docx import Document
import prepare
import reverse
import converter

class ReverseTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp=tempfile.TemporaryDirectory()
        cls.root=Path(cls.temp.name)
        cls.native=cls.root/'일반.pdf';cls.scan=cls.root/'스캔.pdf'
        with pymupdf.open() as doc:
            p=doc.new_page()
            p.insert_text((50,60),'한글 문서 변환 테스트',fontname='korea',fontsize=18)
            p.insert_text((50,100),'한국어 문장을 수정할 수 있습니다.',fontname='korea',fontsize=12)
            for y in (150,190,230):p.draw_line((50,y),(350,y))
            for x in (50,200,350):p.draw_line((x,150),(x,230))
            for x,y,t in [(65,175,'항목'),(215,175,'금액'),(65,215,'교재'),(215,215,'10000')]:p.insert_text((x,y),t,fontname='korea',fontsize=12)
            doc.save(cls.native)
            with pymupdf.open() as raster:
                q=raster.new_page();q.insert_image(q.rect,pixmap=p.get_pixmap(dpi=300));raster.save(cls.scan)
        cls.docx,cls.report=prepare.prepare(cls.native,cls.root/'native-out',log=lambda _:None)
        cls.ocr_docx,cls.ocr_report=prepare.prepare(cls.scan,cls.root/'ocr-out',log=lambda _:None)
    @classmethod
    def tearDownClass(cls):cls.temp.cleanup()

    def test_native_korean_and_table(self):
        self.assertFalse(self.report['pages'][0]['ocr'])
        self.assertIn('한국어 문장을 수정할 수 있습니다.',prepare.docx_text(self.docx))
        tables=Document(self.docx).tables
        self.assertEqual(len(tables),1)
        self.assertEqual(tables[0].cell(1,0).text,'교재')
        self.assertEqual(self.report['docx_text_retention'],1.0)

    def test_real_korean_ocr_and_grid(self):
        self.assertTrue(self.ocr_report['pages'][0]['ocr'])
        text=prepare.normalized(prepare.docx_text(self.ocr_docx))
        self.assertIn('한글문서변환테스트',text)
        self.assertIn('한국어문장을수정할수있습니다',text)
        self.assertIn('교재',text)
        self.assertIn('10000',text)
        self.assertTrue(Document(self.ocr_docx).tables)
        self.assertTrue(self.ocr_report['warnings'])

    def test_mixed_pdf(self):
        mixed=self.root/'mixed.pdf'
        with pymupdf.open() as d:
            for file in (self.native,self.scan):
                with pymupdf.open(file) as source:d.insert_pdf(source)
            d.save(mixed)
        docx,report=prepare.prepare(mixed,self.root/'mixed-out',log=lambda _:None)
        self.assertEqual([p['ocr'] for p in report['pages']],[False,True])
        self.assertEqual(prepare.normalized(prepare.docx_text(docx)).count('한글문서변환테스트'),2)

    def test_bad_and_encrypted(self):
        bad=self.root/'bad.pdf';bad.write_bytes(b'bad')
        with self.assertRaises(Exception):prepare.prepare(bad,self.root/'bad-out')
        encrypted=self.root/'locked.pdf'
        with pymupdf.open() as d:d.new_page();d.save(encrypted,encryption=pymupdf.PDF_ENCRYPT_AES_256,user_pw='pw',owner_pw='owner')
        with self.assertRaises(ValueError):prepare.prepare(encrypted,self.root/'locked-out')

    def test_model_integrity(self):
        folder=prepare.verify_models()
        self.assertTrue((folder/'kor.traineddata').exists())

    def fake_prepare(self,src,folder,mode,log):
        shutil.copy2(self.docx,folder/'편집용.docx')
        (folder/'추출문자.txt').write_text('한글')
        return folder/'편집용.docx',json.loads(json.dumps(self.report))

    def test_completion_skip_and_deleted_hwp_recovery(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp);watch=root/'input';watch.mkdir();src=watch/'보고서.pdf';shutil.copy2(self.native,src)
            engine=MagicMock()
            def save(a,b):
                Path(b).write_bytes(bytes.fromhex('d0cf11e0a1b11ae1')+b'test only')
                return {'hancom_version':'MOCK','hwp_text_retention':1.0}
            engine.convert.side_effect=save
            with patch.object(reverse,'prepare',side_effect=self.fake_prepare):
                self.assertTrue(reverse.process(src,watch,root/'output',engine))
                self.assertFalse(reverse.process(src,watch,root/'output',engine))
                manifest=root/'output'/src.name/'.complete.json'
                old=json.loads(manifest.read_text(encoding='utf-8'));old['version']=1
                manifest.write_text(json.dumps(old),encoding='utf-8')
                self.assertTrue(reverse.process(src,watch,root/'output',engine))
                (root/'output'/(src.name+' (1)')/'보고서.hwp').unlink()
                self.assertTrue(reverse.process(src,watch,root/'output',engine))
            self.assertTrue(src.exists())

    def test_failed_hwp_keeps_docx_not_completed(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp);watch=root/'in';watch.mkdir();src=watch/'보고서.pdf';shutil.copy2(self.native,src)
            engine=MagicMock();engine.convert.side_effect=RuntimeError('HWP failure')
            with patch.object(reverse,'prepare',side_effect=self.fake_prepare):
                with self.assertRaises(RuntimeError):reverse.process(src,watch,root/'out',engine)
            self.assertTrue((root/'확인 필요'/src.name/'편집용.docx').exists())
            self.assertFalse(list((root/'out').rglob('.complete.json')))

    def test_hwp_2018_2022_common_com_contract_mock(self):
        for version in ('10.0 (2018 mock)','12.0 (2022 mock)'):
            with self.subTest(version=version),tempfile.TemporaryDirectory() as tmp:
                dst=Path(tmp)/'result.hwp';app=MagicMock();pc=MagicMock();reg=MagicMock()
                app.Version=version;app.Open.return_value=True;app.RegisterModule.return_value=True
                app.GetTextFile.return_value=prepare.docx_text(self.docx)
                def save(*a):dst.write_bytes(bytes.fromhex('d0cf11e0a1b11ae1')+b'mock');return True
                app.SaveAs.side_effect=save
                client=types.ModuleType('win32com.client');client.DispatchEx=MagicMock(return_value=app)
                package=types.ModuleType('win32com');package.client=client
                with patch.dict(sys.modules,{'win32com':package,'win32com.client':client,'pythoncom':pc,'winreg':reg}):
                    result=converter.save_hwp(self.docx,dst)
                self.assertEqual(result['hwp_text_retention'],1.0)
                app.SaveAs.assert_called_once_with(str(dst),'HWP','')
                self.assertEqual(app.Open.call_count,2)
                self.assertFalse(app.XHwpWindows.Item(0).Visible)
                app.Quit.assert_called_once()



class DocxImportRegressionTests(unittest.TestCase):
    def test_requires_insert_option_to_avoid_hanword(self):
        app=MagicMock()
        app.Open.side_effect=lambda path,fmt,args: 'insert:true' in args
        method=converter.open_docx_in_hwp(app,Path('input.docx'))
        self.assertEqual(method,'OOXML 삽입')
        app.Open.assert_called_once_with('input.docx','OOXML','forceopen:true;insert:true')
        app.Clear.assert_called_once_with(1)

    def test_fallback_keeps_insert_and_clears_partial_document(self):
        app=MagicMock();app.Open.side_effect=[False,True]
        method=converter.open_docx_in_hwp(app,Path('input.docx'))
        self.assertEqual(method,'자동 인식 삽입')
        self.assertEqual(app.Clear.call_count,2)
        self.assertEqual(app.Open.call_args_list[1].args,('input.docx','','forceopen:true;insert:true'))

    def test_both_failed_raise_no_false_success(self):
        app=MagicMock();app.Open.return_value=False
        with self.assertRaisesRegex(RuntimeError,'Open 반환값 False'):
            converter.open_docx_in_hwp(app,Path('input.docx'))
        app.SaveAs.assert_not_called()

if __name__=='__main__':unittest.main()
