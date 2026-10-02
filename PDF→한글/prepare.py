"""PDF 분석/OCR/편집용 DOCX. OCR 결과는 원문으로 보정하거나 추측하지 않는다."""
from collections import Counter
import hashlib
import json
import os
from pathlib import Path
import re
import unicodedata
import pymupdf

BASE = Path(__file__).resolve().parent


def normalized(text):
    return ''.join(c for c in unicodedata.normalize('NFC', text) if c.isalnum())


def text_coverage(source, target):
    a, b = Counter(normalized(source)), Counter(normalized(target))
    return sum((a & b).values()) / sum(a.values()) if a else 1.0


def verify_models():
    folder = BASE / 'tessdata'
    manifest = json.loads((folder / 'manifest.json').read_text(encoding='utf-8'))
    for name in ('kor.traineddata', 'eng.traineddata'):
        data = (folder / name).read_bytes()
        if hashlib.sha256(data).hexdigest() != manifest[name]['sha256']:
            raise RuntimeError('OCR 언어 데이터가 손상되었습니다. ZIP을 다시 압축 해제하세요: ' + name)
    return folder


def needs_ocr(page, mode):
    if mode == 'always':
        return True
    text = page.get_text()
    bad = sum(c == '\ufffd' or 0xE000 <= ord(c) <= 0xF8FF for c in text)
    area = max(1, (page.rect.width * page.rect.height))
    images = page.get_image_info()
    image_area = sum(max(0, r.width) * max(0, r.height) for r in (pymupdf.Rect(im['bbox']) & page.rect for im in images))
    if bad > max(1, len(text) * 0.02):
        return True
    if image_area / area > 0.45:
        return True
    return not text.strip() and bool(images or page.get_drawings())


def grid_lines(page, dpi=150):
    """스캔본의 긴 수평/수직 선만 검출. 복잡한 병합 표의 완전 복원은 보장하지 않음."""
    import cv2
    import numpy as np
    pix = page.get_pixmap(dpi=dpi, colorspace=pymupdf.csGRAY, alpha=False)
    pixels = np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.height, pix.width)
    ink = cv2.threshold(pixels, 0, 255, cv2.THRESH_BINARY_INV + cv2.THRESH_OTSU)[1]
    result = []
    for horizontal in (True, False):
        length = max(35, (pix.width if horizontal else pix.height) // 25)
        kernel = cv2.getStructuringElement(cv2.MORPH_RECT, (length, 1) if horizontal else (1, length))
        mask = cv2.morphologyEx(ink, cv2.MORPH_OPEN, kernel)
        contours, _ = cv2.findContours(mask, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE)
        for contour in contours:
            x, y, w, h = cv2.boundingRect(contour)
            if (horizontal and w >= length and h <= 12) or (not horizontal and h >= length and w <= 12):
                scale = 72 / dpi
                if horizontal:
                    result.append(((x*scale, (y+h/2)*scale), ((x+w)*scale, (y+h/2)*scale)))
                else:
                    result.append((((x+w/2)*scale, y*scale), ((x+w/2)*scale, (y+h)*scale)))
    return result[:2000]


def ocr_page(page, target, tessdata, warnings):
    if (page.rect.width * page.rect.height) * (300 / 72) ** 2 > 60_000_000:
        raise ValueError('OCR 페이지가 너무 큽니다. A4/A3 크기의 사본으로 다시 시도하세요.')
    # 실제 보이는 방향으로 정규화하여 텍스트 좌표와 페이지 크기를 맞춘다.
    with pymupdf.open() as raster:
        image = page.get_pixmap(dpi=300, colorspace=pymupdf.csRGB, alpha=False)
        rp = raster.new_page(width=page.rect.width, height=page.rect.height)
        # 표 선이 문자를 가리지 않도록 OCR용 복사본에서 긴 선만 제거.
        import cv2
        import numpy as np
        pixels=np.frombuffer(image.samples,dtype=np.uint8).reshape(image.height,image.width,3)
        gray=cv2.cvtColor(pixels,cv2.COLOR_RGB2GRAY)
        ink=cv2.threshold(gray,0,255,cv2.THRESH_BINARY_INV+cv2.THRESH_OTSU)[1]
        horizontal=cv2.morphologyEx(ink,cv2.MORPH_OPEN,
            cv2.getStructuringElement(cv2.MORPH_RECT,(max(60,image.width//25),1)))
        vertical=cv2.morphologyEx(ink,cv2.MORPH_OPEN,
            cv2.getStructuringElement(cv2.MORPH_RECT,(1,max(60,image.height//25))))
        clean=gray.copy()
        clean[cv2.bitwise_or(horizontal,vertical)>0]=255
        ok,png=cv2.imencode('.png',clean)
        if not ok: raise RuntimeError('OCR 이미지 준비 실패')
        rp.insert_image(rp.rect,stream=png.tobytes())
        tp = rp.get_textpage_ocr(language=os.environ.get('PDF_OCR_LANGUAGE','kor'), dpi=300, full=True, tessdata=str(tessdata))
        data = rp.get_text('dict', textpage=tp, sort=True)
        text = rp.get_text(textpage=tp, sort=True)
        if not text.strip():
            target.insert_pdf(page.parent, from_page=page.number, to_page=page.number)
            warnings.append('OCR에서 글자를 찾지 못해 원본 페이지 모양을 유지했습니다. 편집 가능한 글자는 없습니다.')
            return text, 0
        dest = target.new_page(width=rp.rect.width, height=rp.rect.height)
        font = pymupdf.Font('korea')
        for block in data['blocks']:
            for line in block.get('lines', []):
                for span in line['spans']:
                    value = unicodedata.normalize('NFC', span['text']).replace('\x00', '')
                    if not value.strip():
                        continue
                    x0, y0, x1, y1 = span['bbox']
                    size = min(max(4, span['size']), max(4, (y1-y0)*0.85))
                    advance = font.text_length(value, fontsize=size)
                    if advance > 0:
                        size *= min(1, max(1, x1-x0)/advance)
                    baseline = min(y1, y0 + (y1-y0)*0.85)
                    dest.insert_text((max(0,x0), baseline), value, fontname='korea', fontsize=max(2,size))
        lines = grid_lines(page)
        for start, end in lines:
            dest.draw_line(start, end, color=(0,0,0), width=0.5)
        warnings.append('OCR 적용: 오인식·띄어쓰기·표 병합을 확인하세요. 사진/도장/손글씨/수식의 모양은 원본 PDF에서 확인하세요.')
        if lines:
            warnings.append('스캔 표의 선을 추정했습니다. 셀 구조가 원본과 다를 수 있습니다.')
        return text, len(lines)


def docx_text(path):
    from docx import Document
    doc = Document(path)
    texts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            texts.extend(cell.text for cell in row.cells)
    return '\n'.join(texts)


def prepare(pdf, output, mode='auto', log=print):
    from pdf2docx import Converter
    from docx_compat import normalize_docx
    output = Path(output)
    output.mkdir(parents=True, exist_ok=True)
    normalized_pdf = output / '.editable.pdf'
    docx_path = output / '편집용.docx'
    report = {'pages': [], 'warnings': [], 'ocr_mode': mode, 'ocr_language':os.environ.get('PDF_OCR_LANGUAGE','kor')}
    expected = []
    tessdata = None
    with pymupdf.open(pdf) as source, pymupdf.open() as target:
        if not source.is_pdf or source.needs_pass or not len(source):
            raise ValueError('암호가 있거나 읽을 수 없는 PDF입니다.')
        for i, page in enumerate(source):
            notes = []
            use_ocr = needs_ocr(page, mode)
            log(f'{i+1}/{len(source)}쪽: ' + ('한국어 OCR' if use_ocr else '원문 글자/표 추출'))
            if use_ocr:
                if tessdata is None:
                    tessdata = verify_models()
                text, lines = ocr_page(page, target, tessdata, notes)
            else:
                target.insert_pdf(source, from_page=i, to_page=i)
                text = page.get_text(sort=True)
                lines = 0
            expected.append(text)
            report['pages'].append({'page': i+1, 'ocr':use_ocr, 'characters':len(text), 'grid_lines':lines, 'warnings':notes})
        target.save(normalized_pdf)
    engine = Converter(str(normalized_pdf))
    try:
        engine.convert(str(docx_path), multi_processing=False, parse_stream_table=False,
                       ignore_page_error=False, raw_exceptions=True,
                       line_break_width_ratio=1.0)
    finally:
        engine.close()
    report['layout_repairs'] = normalize_docx(docx_path)
    if report['layout_repairs']['tables_unresolved']:
        report['warnings'].append('일부 표의 열 너비를 확정하지 못했습니다. 원본 표와 비교하세요.')
    extracted = '\n\n'.join(expected)
    (output / '추출문자.txt').write_text(extracted, encoding='utf-8-sig')
    report['docx_text_retention'] = round(text_coverage(extracted, docx_text(docx_path)),4)
    if report['docx_text_retention'] < 0.97:
        report['warnings'].append('추출 문자 일부가 편집용 문서에서 누락되었을 수 있습니다. 추출문자.txt와 비교하세요.')
    if any(p['ocr'] for p in report['pages']):
        report['warnings'].append('OCR 결과는 원본 대조가 필요합니다. 문자 보존 비율은 OCR 정확도를 의미하지 않습니다.')
    report['warnings'].append('PDF에서 복원한 문서이므로 줄바꿈/쪽수/표/수식/글꼴이 달라질 수 있습니다.')
    normalized_pdf.unlink(missing_ok=True)
    return docx_path, report
